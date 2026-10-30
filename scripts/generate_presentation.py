import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)       # #0f172a
    BG_LIGHT = RGBColor(248, 250, 252)   # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    TEXT_MAIN = RGBColor(30, 41, 59)     # #1e293b
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748b
    BRAND_BLUE = RGBColor(37, 99, 235)   # #2563eb
    BRAND_DARK_BLUE = RGBColor(30, 58, 138) # #1e3a8a
    ACCENT_GREEN = RGBColor(16, 185, 129)# #10b981
    ACCENT_AMBER = RGBColor(245, 158, 11)# #f59e0b
    ACCENT_PURPLE = RGBColor(139, 92, 246)# #8b5cf6
    ACCENT_RED = RGBColor(239, 68, 68)   # #ef4444
    BORDER_LIGHT = RGBColor(226, 232, 240) # #e2e8f0

    def add_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title, subtitle=None):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = BRAND_DARK_BLUE
        p.font.name = "Segoe UI"
        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_MUTED
            p2.font.name = "Segoe UI"
            p2.space_before = Pt(4)

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, BG_DARK)
    
    # Decorative accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.4), Inches(0.8), Inches(0.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_BLUE
    bar.line.fill.background()
    bar_tf = bar.text_frame
    bar_p = bar_tf.paragraphs[0]
    bar_p.text = "🍽️"
    bar_p.font.size = Pt(28)
    bar_p.alignment = PP_ALIGN.CENTER

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(10.5), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Digital Menu &\nOrder Management System\nfor Hotel"
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    p.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Modern. Fast. Organized."
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT_AMBER
    p2.font.bold = True
    p2.font.name = "Segoe UI"
    p2.space_after = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Better Service  •  Happier Customers  •  Higher Efficiency"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 2: PURPOSE / USE OF SYSTEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, BG_LIGHT)
    add_header(s2, "1. What is the purpose / Use of system?", 
               "The Digital Menu and Order Management System is designed to make hotel and restaurant operations faster, easier and more organized.")

    roles = [
        ("👤 Customers", "Scan QR code, view menu and place orders.", ACCENT_GREEN),
        ("🛎️ Waiters", "Receive orders, monitor kitchen, and serve customers.", BRAND_BLUE),
        ("💵 Cashier", "Manage payments, assign waiters, and financial records.", ACCENT_AMBER),
        ("👨‍🍳 Kitchen/Chef", "Receive orders, prepare food, and update status.", ACCENT_PURPLE),
        ("👔 Admin/Manager", "Control the system, manage staff, approve changes, and view reports.", BRAND_DARK_BLUE),
    ]
    top_pos = 1.9
    for title, desc, col in roles:
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.7), Inches(0.9))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        box.line.width = Pt(1)
        
        # Pill
        pill = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(top_pos + 0.15), Inches(2.6), Inches(0.6))
        pill.fill.solid()
        pill.fill.fore_color.rgb = col
        pill.line.fill.background()
        ptf = pill.text_frame
        pp = ptf.paragraphs[0]
        pp.text = title
        pp.font.size = Pt(13)
        pp.font.bold = True
        pp.font.color.rgb = WHITE
        pp.alignment = PP_ALIGN.CENTER
        
        # Text
        dtb = s2.shapes.add_textbox(Inches(3.9), Inches(top_pos + 0.15), Inches(8.3), Inches(0.6))
        dtf = dtb.text_frame
        dp = dtf.paragraphs[0]
        dp.text = desc
        dp.font.size = Pt(13)
        dp.font.color.rgb = TEXT_MAIN
        dp.font.name = "Segoe UI"
        top_pos += 1.02

    # ==========================================
    # SLIDE 3: HOW DOES IT WORK?
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, BG_LIGHT)
    add_header(s3, "2. How does it work?", "Step-by-step workflow from guest arrival to order settlement and reports.")

    steps_top = [
        ("1. Scan QR Code", "Customer scans QR code on the table.", BRAND_BLUE),
        ("2. View Menu", "Customer browses menu and selects dishes.", ACCENT_GREEN),
        ("3. Place Order", "Order is sent directly to the table waiter.", BRAND_DARK_BLUE),
        ("4. Waiter Receives", "Waiter checks order and forwards to kitchen.", ACCENT_PURPLE),
    ]
    steps_bot = [
        ("5. Kitchen Accepts", "Chef accepts order and starts cooking timer.", ACCENT_AMBER),
        ("6. Finished", "Chef marks 'Finished' when food is plated.", ACCENT_GREEN),
        ("7. Served", "Waiter serves food to table and marks 'Served'.", BRAND_BLUE),
        ("8. Payment", "Customer pays waiter; Cashier settles bill.", ACCENT_PURPLE),
        ("9. Reports", "Admin and cashier can download reports.", BRAND_DARK_BLUE),
    ]

    # Row 1
    left = 0.8
    for st, sd, sc in steps_top:
        b = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.9), Inches(2.75), Inches(2.2))
        b.fill.solid()
        b.fill.fore_color.rgb = WHITE
        b.line.color.rgb = BORDER_LIGHT
        
        # header bar inside card
        hb = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.15), Inches(2.05), Inches(2.45), Inches(0.5))
        hb.fill.solid()
        hb.fill.fore_color.rgb = sc
        hb.line.fill.background()
        hbp = hb.text_frame.paragraphs[0]
        hbp.text = st
        hbp.font.size = Pt(12)
        hbp.font.bold = True
        hbp.font.color.rgb = WHITE
        hbp.alignment = PP_ALIGN.CENTER
        
        tb = s3.shapes.add_textbox(Inches(left + 0.15), Inches(2.7), Inches(2.45), Inches(1.2))
        tbp = tb.text_frame.paragraphs[0]
        tbp.text = sd
        tbp.font.size = Pt(11)
        tbp.font.color.rgb = TEXT_MAIN
        tbp.font.name = "Segoe UI"
        left += 2.98

    # Row 2
    left = 0.8
    card_w = 2.18
    for st, sd, sc in steps_bot:
        b = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(4.5), Inches(card_w), Inches(2.2))
        b.fill.solid()
        b.fill.fore_color.rgb = WHITE
        b.line.color.rgb = BORDER_LIGHT
        
        hb = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left + 0.12), Inches(4.65), Inches(card_w - 0.24), Inches(0.5))
        hb.fill.solid()
        hb.fill.fore_color.rgb = sc
        hb.line.fill.background()
        hbp = hb.text_frame.paragraphs[0]
        hbp.text = st
        hbp.font.size = Pt(11)
        hbp.font.bold = True
        hbp.font.color.rgb = WHITE
        hbp.alignment = PP_ALIGN.CENTER
        
        tb = s3.shapes.add_textbox(Inches(left + 0.12), Inches(5.3), Inches(card_w - 0.24), Inches(1.2))
        tbp = tb.text_frame.paragraphs[0]
        tbp.text = sd
        tbp.font.size = Pt(10)
        tbp.font.color.rgb = TEXT_MAIN
        tbp.font.name = "Segoe UI"
        left += card_w + 0.2

    # ==========================================
    # SLIDE 4: REQUIREMENTS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, BG_LIGHT)
    add_header(s4, "3. Requirements", "The hotel needs the following basic equipment (no expensive POS servers needed):")

    reqs = [
        ("🌐 Internet Connection", "Basic Wi-Fi or 4G router (uses < 300 KB per order)"),
        ("🏷️ Table QR Codes", "Printed acrylic QR code stands for each dining table"),
        ("📱 Smartphones / Tablets", "Budget tablets for kitchen display and waiters"),
        ("💻 Computer / Laptop", "Any desktop or laptop for cashier and manager"),
        ("🖨️ Thermal Printer (Optional)", "Standard 80mm or 58mm printer for guest receipts"),
        ("🍽️ Food Information", "Dish names, descriptions, prices, photos, and tags"),
        ("👥 Employee Information", "Staff names, roles, and login emails for staff accounts"),
    ]
    grid_x = [0.8, 6.8]
    grid_y = [1.9, 3.1, 4.3, 5.5]
    for idx, (title, desc) in enumerate(reqs):
        x = grid_x[idx % 2]
        y = grid_y[idx // 2]
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        
        tb = s4.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.12), Inches(5.3), Inches(0.76))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = BRAND_DARK_BLUE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # Bottom highlight
    callout = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(5.5), Inches(5.7), Inches(1.0))
    callout.fill.solid()
    callout.fill.fore_color.rgb = BRAND_BLUE
    callout.line.fill.background()
    ctf = callout.text_frame
    cp = ctf.paragraphs[0]
    cp.text = "💡 Zero Customer App Installation"
    cp.font.size = Pt(13)
    cp.font.bold = True
    cp.font.color.rgb = WHITE
    cp2 = ctf.add_paragraph()
    cp2.text = "Customers access the menu instantly in their mobile browser (Safari / Chrome)."
    cp2.font.size = Pt(11)
    cp2.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 5: WHEN SYSTEM LAG OR CRASH?
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, BG_LIGHT)
    add_header(s5, "4. When the system lag or crash?", "We have a built-in plan to keep your hotel running smoothly with zero downtime.")

    # Left Column: Problems
    p_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.7), Inches(3.8))
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = WHITE
    p_box.line.color.rgb = RGBColor(254, 202, 202)
    
    ptb = s5.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.1), Inches(3.4))
    ptf = ptb.text_frame
    pp = ptf.paragraphs[0]
    pp.text = "⚠️ Possible Problems"
    pp.font.size = Pt(15)
    pp.font.bold = True
    pp.font.color.rgb = ACCENT_RED
    
    for item in ["• Slow internet connection", "• System becomes slow during peak hours", "• System temporarily stops", "• Employee cannot log in or forgets password", "• Order not displayed on screen", "• Hardware device battery or screen problem"]:
        par = ptf.add_paragraph()
        par.text = item
        par.font.size = Pt(12)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(6)

    # Right Column: Solutions
    s_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(3.8))
    s_box.fill.solid()
    s_box.fill.fore_color.rgb = WHITE
    s_box.line.color.rgb = RGBColor(167, 243, 208)
    
    stb = s5.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(3.4))
    stf = stb.text_frame
    sp = stf.paragraphs[0]
    sp.text = "✅ Our Built-in Solution"
    sp.font.size = Pt(15)
    sp.font.bold = True
    sp.font.color.rgb = ACCENT_GREEN
    
    for item in ["• Offline cache stores orders and syncs when reconnected", "• Cloud database handles 1,000+ orders/day with 0.2s sync", "• Reopening tab restores tickets & timers immediately", "• Manager resets staff passwords with 1-click in /admin", "• Cashier checks connection & order status in real time", "• Switch to another authorized phone or tablet in seconds"]:
        par = stf.add_paragraph()
        par.text = item
        par.font.size = Pt(12)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(6)

    # Bottom Callout
    bot_callout = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9))
    bot_callout.fill.solid()
    bot_callout.fill.fore_color.rgb = BRAND_DARK_BLUE
    bot_callout.line.fill.background()
    bctf = bot_callout.text_frame
    bcp = bctf.paragraphs[0]
    bcp.text = "🛡️ Data Protection Guarantee"
    bcp.font.size = Pt(13)
    bcp.font.bold = True
    bcp.font.color.rgb = WHITE
    bcp2 = bctf.add_paragraph()
    bcp2.text = "Orders and financial records are stored in a secure cloud database with daily backups. All records remain 100% safe."
    bcp2.font.size = Pt(11)
    bcp2.font.color.rgb = RGBColor(226, 232, 240)

    # ==========================================
    # SLIDE 6: DOES SYSTEM HAVE EXPIRY DATE?
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, BG_LIGHT)
    add_header(s6, "5. Does the system have an expiry date?", "Transparent breakdown of system ownership, hosting, and technical support.")

    cards = [
        ("🏛️ System Ownership", "Permanent Usage Rights", [
            "• Hotel owns full usage rights to the system.",
            "• All database records and history belong 100% to the hotel.",
            "• 1-click offline JSON backups available anytime in /admin.",
            "• No artificial software lockouts or sudden expiration.",
        ], BRAND_DARK_BLUE),
        ("🌐 Hosting & Domain", "Annual Domain Renewal", [
            "• Web domain (e.g. hiragahotel.com) renewed annually.",
            "• Cloud Database runs permanently on Supabase.",
            "• 100% Free plan supports up to 500,000 lifetime orders.",
            "• Optional Pro plan ($25/mo) for multi-year photo archiving.",
        ], BRAND_BLUE),
        ("🛠️ Technical Support", "Maintenance Agreement", [
            "• Launch monitoring during opening week peak hours.",
            "• Free warranty and bug-fix period per agreement.",
            "• Manager has self-service controls for prices & staff.",
            "• Continuous technical support via phone & Telegram.",
        ], ACCENT_GREEN),
    ]
    c_left = 0.8
    for icon_title, sub, bullets, c_col in cards:
        cb = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(1.9), Inches(3.7), Inches(4.3))
        cb.fill.solid()
        cb.fill.fore_color.rgb = WHITE
        cb.line.color.rgb = BORDER_LIGHT
        
        # header
        ch = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left + 0.15), Inches(2.05), Inches(3.4), Inches(0.85))
        ch.fill.solid()
        ch.fill.fore_color.rgb = c_col
        ch.line.fill.background()
        ch_tf = ch.text_frame
        cp1 = ch_tf.paragraphs[0]
        cp1.text = icon_title
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = WHITE
        cp1.alignment = PP_ALIGN.CENTER
        cp2 = ch_tf.add_paragraph()
        cp2.text = sub
        cp2.font.size = Pt(10)
        cp2.font.color.rgb = RGBColor(226, 232, 240)
        cp2.alignment = PP_ALIGN.CENTER
        
        # bullets
        ctb = s6.shapes.add_textbox(Inches(c_left + 0.2), Inches(3.1), Inches(3.3), Inches(2.9))
        ctf = ctb.text_frame
        for b_idx, b_text in enumerate(bullets):
            bp = ctf.paragraphs[0] if b_idx == 0 else ctf.add_paragraph()
            bp.text = b_text
            bp.font.size = Pt(11)
            bp.font.color.rgb = TEXT_MAIN
            bp.space_before = Pt(8)
        
        c_left += 4.0

    # ==========================================
    # SLIDE 7: OUR INSTALLATION
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7, BG_LIGHT)
    add_header(s7, "6. Our installation process", "We will set up the entire hotel system step-by-step in 8 structured phases:")

    inst_steps = [
        ("1. Collect Requirements", "Review hotel layout, table count, and payment methods"),
        ("2. Register Hotel & Cloud", "Configure database, secure rules, and serverless hosting"),
        ("3. Add Menu (Food & Drinks)", "Input dishes, descriptions, Amharic translations & photos"),
        ("4. Create Staff Accounts", "Set up accounts for Cashiers, Chefs, and Waiters"),
        ("5. Generate Table QR Codes", "Produce high-resolution QR table cards for every table"),
        ("6. Connect & Test Devices", "Pair Cashier PC, Kitchen tablet, and Waiter devices"),
        ("7. Train Hotel Employees", "Hands-on 30-minute practice simulation for all staff"),
        ("8. Final Testing & Launch", "Simulated dinner rush test before opening to public"),
    ]
    for idx, (st, sd) in enumerate(inst_steps):
        x = 0.8 if idx < 4 else 6.8
        y = 1.9 + (idx % 4) * 1.2
        box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(1.05))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        
        tb = s7.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.12), Inches(5.3), Inches(0.8))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = st
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = BRAND_DARK_BLUE
        p2 = tf.add_paragraph()
        p2.text = sd
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 8: MENU FEATURES
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, BG_LIGHT)
    add_header(s8, "7. Digital Menu & Features", "The digital menu is what customers experience when scanning the table QR code.")

    # Customer view features
    m_box1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.7), Inches(4.8))
    m_box1.fill.solid()
    m_box1.fill.fore_color.rgb = WHITE
    m_box1.line.color.rgb = BORDER_LIGHT
    
    tb1 = s8.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.1), Inches(4.3))
    tf1 = tb1.text_frame
    p = tf1.paragraphs[0]
    p.text = "📱 Customer Digital Menu Includes:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    for item in [
        "✅ Clean food and drink categories with instant search",
        "✅ High-resolution pictures and bilingual descriptions",
        "✅ Clear pricing in ETB (Ethiopian Birr)",
        "✅ Instant Available / Sold-out status indicators",
        "✅ Dietary tags: Fasting (የጾም), Special, Spicy (የሚቃጥል)",
        "✅ Custom preparation notes (e.g. 'ያለ ሽንኩርት / No onions')",
        "✅ Digital shopping cart with quantity adjustments",
        "✅ Hypermedia icons (Phone, Telegram, WhatsApp, Maps)",
    ]:
        par = tf1.add_paragraph()
        par.text = item
        par.font.size = Pt(11.5)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(6)

    # Admin controls
    m_box2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.8))
    m_box2.fill.solid()
    m_box2.fill.fore_color.rgb = WHITE
    m_box2.line.color.rgb = BORDER_LIGHT
    
    tb2 = s8.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.3))
    tf2 = tb2.text_frame
    p = tf2.paragraphs[0]
    p.text = "⚙️ Hotel Manager Can:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    for item in [
        "➕ Add new dishes or drinks in 10 seconds in /admin",
        "✏️ Update food prices and photos with instant live sync",
        "📁 Create and re-order menu categories",
        "🔴 Toggle dishes 'Sold Out' or 'Available' with 1 click",
        "⭐ Mark 'Today's Special' to highlight featured dishes",
        "💳 Update Telebirr and CBE Bank payment account details",
        "👥 Create, edit, and deactivate waiter and chef accounts",
        "🔑 Reset staff passwords and unlock accounts directly",
    ]:
        par = tf2.add_paragraph()
        par.text = item
        par.font.size = Pt(11.5)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(6)

    # ==========================================
    # SLIDE 9: USE OF MENU (CUSTOMER STEPS)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9, BG_LIGHT)
    add_header(s9, "8. Use of menu (Customer Experience)", "Simple, frictionless 4-step experience for dining guests:")

    cust_steps = [
        ("Step 1: Scan QR Code", "Guest sits at table and points phone camera at the acrylic QR card. Opens instantly in browser."),
        ("Step 2: Browse Menu", "Explore dishes by category, view vivid photos, ingredients, prices, and dietary tags."),
        ("Step 3: Add to Cart", "Select quantity and type custom kitchen requests (e.g. 'mild / no pepper', 'extra sauce')."),
        ("Step 4: Send to Waiter", "Order is placed and sent to table waiter. Food is served, and guest pays after dining."),
    ]
    cs_left = 0.8
    for idx, (st, sd) in enumerate(cust_steps):
        box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cs_left), Inches(1.9), Inches(2.75), Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        
        # Number badge
        nb = s9.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cs_left + 0.95), Inches(2.2), Inches(0.85), Inches(0.85))
        nb.fill.solid()
        nb.fill.fore_color.rgb = BRAND_BLUE if idx % 2 == 0 else ACCENT_AMBER
        nb.line.fill.background()
        nb_p = nb.text_frame.paragraphs[0]
        nb_p.text = str(idx + 1)
        nb_p.font.size = Pt(20)
        nb_p.font.bold = True
        nb_p.font.color.rgb = WHITE
        nb_p.alignment = PP_ALIGN.CENTER
        
        tb = s9.shapes.add_textbox(Inches(cs_left + 0.15), Inches(3.3), Inches(2.45), Inches(2.5))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = st
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = BRAND_DARK_BLUE
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = sd
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(8)
        
        cs_left += 2.98

    # ==========================================
    # SLIDE 10: WHAT DOES EACH EMPLOYEE SEE?
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, BG_LIGHT)
    add_header(s10, "9. What does each employee see?", "Dedicated, role-specific interfaces designed for each station.")

    emp_cards = [
        ("👔 Admin Dashboard (/admin)", BRAND_DARK_BLUE, [
            "• Comprehensive hotel activity overview",
            "• Manage menu dishes, prices, and categories",
            "• Staff account creation and password resets",
            "• Downloadable sales and order reports",
            "• Bank accounts & social media management",
        ]),
        ("🛎️ Waiter Dashboard (/waiter)", BRAND_BLUE, [
            "• View customer table order requests",
            "• Monitor active cooking progress in kitchen",
            "• Instant chime when kitchen marks food ready",
            "• Mark orders 'Delivered to Table'",
            "• Take mobile orders with table POS",
        ]),
        ("💵 Cashier Dashboard (/cashier)", ACCENT_AMBER, [
            "• View incoming orders and pending bills",
            "• Tag and assign Waiter Names to orders",
            "• Filter orders by specific waiter (no ambiguity)",
            "• Settle post-served customer bills",
            "• Print thermal 80mm/58mm receipts",
        ]),
        ("👨‍🍳 Kitchen KDS (/kitchen)", ACCENT_PURPLE, [
            "• Receive tickets with prominent table numbers",
            "• Live elapsed cooking timers (Green/Yellow/Red)",
            "• Clear display of assigned server/waiter",
            "• Strike-through ingredient checklists",
            "• 1-Click 'Plated & Ready' pickup alert",
        ]),
    ]
    for idx, (et, ec, elist) in enumerate(emp_cards):
        x = 0.8 if idx % 2 == 0 else 6.8
        y = 1.9 if idx < 2 else 4.4
        box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        
        # header tag
        ht = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.15), Inches(y + 0.15), Inches(5.4), Inches(0.45))
        ht.fill.solid()
        ht.fill.fore_color.rgb = ec
        ht.line.fill.background()
        hp = ht.text_frame.paragraphs[0]
        hp.text = et
        hp.font.size = Pt(12)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER
        
        tb = s10.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.65), Inches(5.3), Inches(1.5))
        tf = tb.text_frame
        for l_idx, line in enumerate(elist):
            lp = tf.paragraphs[0] if l_idx == 0 else tf.add_paragraph()
            lp.text = line
            lp.font.size = Pt(10.5)
            lp.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 11: WHEN HOTEL NEEDS A CHANGE
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_bg(s11, BG_LIGHT)
    add_header(s11, "10. When the hotel needs a change", "You can request changes anytime. We follow a clear, collaborative process.")

    # Left
    b1 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.7), Inches(4.8))
    b1.fill.solid()
    b1.fill.fore_color.rgb = WHITE
    b1.line.color.rgb = BORDER_LIGHT
    
    t1 = s11.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.1), Inches(4.3))
    p = t1.text_frame.paragraphs[0]
    p.text = "📋 Common Hotel Changes"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BRAND_DARK_BLUE
    for item in [
        "• Change food or beverage prices",
        "• Add seasonal dishes or drinks",
        "• Update food descriptions or food photos",
        "• Remove or hide discontinued menu items",
        "• Add new staff accounts (Cashier, Chef, Waiter)",
        "• Adjust permissions and station access",
        "• Add new tables or new dining areas",
        "• Custom reports or accounting exports",
    ]:
        par = t1.text_frame.add_paragraph()
        par.text = item
        par.font.size = Pt(12)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(6)

    # Right
    b2 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.8))
    b2.fill.solid()
    b2.fill.fore_color.rgb = WHITE
    b2.line.color.rgb = BORDER_LIGHT
    
    t2 = s11.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.3))
    p = t2.text_frame.paragraphs[0]
    p.text = "🔄 Our Streamlined Process"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    for item in [
        "1. You inform us about the desired update",
        "2. Immediate guidance: (Many changes like price updates can be done instantly by the manager in /admin)",
        "3. For custom features, we discuss and define requirements",
        "4. We agree on turnaround time (often same-day)",
        "5. We implement and thoroughly test in staging",
        "6. We deploy the update live with zero downtime",
    ]:
        par = t2.text_frame.add_paragraph()
        par.text = item
        par.font.size = Pt(12)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(8)

    # ==========================================
    # SLIDE 12: HOW DOES DIGITAL MENU WORK?
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_bg(s12, BG_LIGHT)
    add_header(s12, "11. How does the digital menu work?", "Core elements engineered for maximum guest conversion and upselling.")

    pillars = [
        ("📁 Categories", "Clear food groupings (Traditional, Fast Foods, Drinks, Desserts) for effortless navigation.", BRAND_BLUE),
        ("📸 Food Pictures", "Crisp photos that drive appetite and higher average order values.", ACCENT_GREEN),
        ("💵 Clear Pricing", "Transparent prices in ETB with zero surprise charges for diners.", ACCENT_AMBER),
        ("🔢 Quantity Controls", "Simple '+' and '-' buttons to easily adjust servings and add extras.", ACCENT_PURPLE),
        ("🛒 Shopping Cart", "Floating order drawer showing live total and item summary.", BRAND_DARK_BLUE),
        ("⚡ Fast Ordering", "1-Tap order submission directly to the assigned table waiter.", ACCENT_GREEN),
    ]
    for idx, (pt, pd, pc) in enumerate(pillars):
        x = 0.8 if idx % 2 == 0 else 6.8
        y = 1.9 + (idx // 2) * 1.6
        box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(1.35))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        
        tb = s12.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.15), Inches(5.2), Inches(1.05))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = pt
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = pc
        p2 = tf.add_paragraph()
        p2.text = pd
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_before = Pt(4)

    # ==========================================
    # SLIDE 13: REPORTS & RECORD PROTECTION
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_bg(s13, BG_LIGHT)
    add_header(s13, "12. Reports and records protection", "Your hotel's financial data and order history is safe, auditable, and always accessible.")

    r_box = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.2))
    r_box.fill.solid()
    r_box.fill.fore_color.rgb = WHITE
    r_box.line.color.rgb = BORDER_LIGHT
    
    rtb = s13.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(3.6))
    rtf = rtb.text_frame
    p = rtf.paragraphs[0]
    p.text = "📊 Key Reporting & Accountability Rules:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BRAND_DARK_BLUE
    for item in [
        "✓ Hotel Manager can download daily, weekly, and monthly sales reports",
        "✓ Cashier can generate verified shift reports tagged by cashier and waiter name",
        "✓ Reports include total revenue, payment breakdown (Cash vs Telebirr vs CBE), and dish popularity",
        "✓ Employees cannot tamper with or delete historical orders from the system",
        "✓ Complete historical audit log maintains accountability across shifts",
        "✓ 1-Click complete JSON system backup downloadable anytime in /admin",
    ]:
        par = rtf.add_paragraph()
        par.text = item
        par.font.size = Pt(12.5)
        par.font.color.rgb = TEXT_MAIN
        par.space_before = Pt(8)

    # Bottom banner
    b_banner = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    b_banner.fill.solid()
    b_banner.fill.fore_color.rgb = BRAND_BLUE
    b_banner.line.fill.background()
    bbp = b_banner.text_frame.paragraphs[0]
    bbp.text = "🔒 This ensures 100% financial transparency and prevents revenue leaks."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = WHITE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 14: IMPORTANT QUESTIONS TO CONFIRM
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_bg(s14, BG_LIGHT)
    add_header(s14, "13. Important questions to confirm before installation", "We align on these 10 operational points to tailor the setup perfectly for your hotel:")

    qs = [
        "1. Do customers pay after dining (recommended) or upfront?",
        "2. How many cashier desks operate simultaneously?",
        "3. How many waiters and kitchen staff will use the system?",
        "4. Will each table have a permanent unique acrylic QR stand?",
        "5. Does the cashier desk need printed thermal receipts?",
        "6. Do you prefer daily or weekly automated sales reports?",
        "7. Who is the designated manager authorized to approve price changes?",
        "8. What is the primary hotel internet connection (Wi-Fi or 4G)?",
        "9. What are the hotel's official Telebirr and CBE account numbers?",
        "10. Who will be our primary technical liaison for day-to-day operations?",
    ]
    for idx, q_text in enumerate(qs):
        x = 0.8 if idx < 5 else 6.8
        y = 1.9 + (idx % 5) * 0.95
        q_box = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(0.8))
        q_box.fill.solid()
        q_box.fill.fore_color.rgb = WHITE
        q_box.line.color.rgb = BORDER_LIGHT
        
        tb = s14.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.1), Inches(5.4), Inches(0.6))
        p = tb.text_frame.paragraphs[0]
        p.text = q_text
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = BRAND_DARK_BLUE

    # ==========================================
    # SLIDE 15: SUMMARY
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    add_bg(s15, BG_LIGHT)
    add_header(s15, "14. Summary & Value Proposition", "The Digital Menu & Order Management System delivers measurable results:")

    benefits = [
        ("🌟 Improve Guest Experience", "Fast digital ordering, vivid food photos, and zero waiting for menus.", BRAND_BLUE),
        ("⚡ Increase Staff Efficiency", "Waiters spend less time writing orders and more time attending to guests.", ACCENT_GREEN),
        ("📋 Organize Operations", "Seamless digital handoffs from Table ➔ Waiter ➔ Kitchen ➔ Cashier.", BRAND_DARK_BLUE),
        ("💵 Track Orders & Payments", "Eliminates waiter ambiguity, unpaid checks, and revenue leakage.", ACCENT_AMBER),
        ("📈 Generate Actionable Reports", "Gain instant visibility into top-selling dishes and peak dining hours.", ACCENT_PURPLE),
        ("🚀 Grow Your Business", "Professional image that impresses local diners and international tourists.", BRAND_BLUE),
    ]
    for idx, (bt, bd, bc) in enumerate(benefits):
        x = 0.8 if idx % 2 == 0 else 6.8
        y = 1.9 + (idx // 2) * 1.5
        box = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        
        tb = s15.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.15), Inches(5.2), Inches(1.0))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = bt
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = bc
        p2 = tf.add_paragraph()
        p2.text = bd
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_before = Pt(4)

    # Bottom footer
    s_footer = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.65))
    s_footer.fill.solid()
    s_footer.fill.fore_color.rgb = BRAND_DARK_BLUE
    s_footer.line.fill.background()
    sfp = s_footer.text_frame.paragraphs[0]
    sfp.text = "Together, we build a smarter, faster, and more profitable hotel."
    sfp.font.size = Pt(12)
    sfp.font.bold = True
    sfp.font.color.rgb = WHITE
    sfp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 16: THANK YOU
    # ==========================================
    s16 = prs.slides.add_slide(blank_layout)
    add_bg(s16, BG_DARK)
    
    t_box = s16.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.2))
    ttf = t_box.text_frame
    tp1 = ttf.paragraphs[0]
    tp1.text = "Thank You!"
    tp1.font.size = Pt(48)
    tp1.font.bold = True
    tp1.font.color.rgb = WHITE
    tp1.alignment = PP_ALIGN.CENTER
    
    tp2 = ttf.add_paragraph()
    tp2.text = "We are ready to answer your questions and make this system work for your hotel."
    tp2.font.size = Pt(18)
    tp2.font.color.rgb = RGBColor(203, 213, 225)
    tp2.alignment = PP_ALIGN.CENTER
    tp2.space_before = Pt(14)
    
    tp3 = ttf.add_paragraph()
    tp3.text = "Better Service   •   Better Management   •   A Brighter Future"
    tp3.font.size = Pt(14)
    tp3.font.bold = True
    tp3.font.color.rgb = ACCENT_AMBER
    tp3.alignment = PP_ALIGN.CENTER
    tp3.space_before = Pt(18)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\my\Downloads\digital-hotel"
    out_file = os.path.join(out_dir, "Digital_Menu_and_Order_Management_System_for_Hotel.pptx")
    create_deck(out_file)
